import logging
from typing import Any

# Prefer import pymupdf, fallback to fitz
try:
    import pymupdf as fitz
except ImportError:
    import fitz

logger = logging.getLogger(__name__)


class PDFParser:
    """Robust PDF text extractor and document validator using PyMuPDF."""

    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str = "document.pdf") -> dict[str, Any]:
        """
        Extract text, page count, and assess health of a PDF from raw bytes.
        Returns dictionary containing:
        - success: bool
        - text: str
        - page_count: int
        - is_scanned: bool
        - is_encrypted: bool
        - is_corrupt: bool
        - error: Optional[str]
        """
        result = {
            "success": False,
            "text": "",
            "page_count": 0,
            "is_scanned": False,
            "is_encrypted": False,
            "is_corrupt": False,
            "error": None,
        }

        if not file_bytes or len(file_bytes) == 0:
            result["error"] = "Empty file (0 bytes)."
            result["is_corrupt"] = True
            return result

        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
        except Exception as e:
            result["is_corrupt"] = True
            result["error"] = f"Corrupt or invalid PDF format: {e!s}"
            return result

        try:
            if doc.is_encrypted:
                result["is_encrypted"] = True
                result["error"] = "PDF document is encrypted/password-protected."
                doc.close()
                return result

            result["page_count"] = len(doc)
            if result["page_count"] == 0:
                result["error"] = "PDF has 0 pages."
                result["is_corrupt"] = True
                doc.close()
                return result

            full_text = []
            total_images = 0

            for page_index in range(len(doc)):
                page = doc[page_index]
                text = page.get_text()
                if text:
                    full_text.append(text)
                images = page.get_images()
                total_images += len(images)

            doc.close()

            combined_text = "\n".join(full_text).strip()
            result["text"] = combined_text

            # Detect scanned PDF: little to no selectable text but images present
            if len(combined_text) < 50 and total_images > 0:
                result["is_scanned"] = True
                # Optional fallback or note that OCR is required
                result["text"] = f"[SCANNED DOCUMENT: Contains {total_images} image elements with insufficient direct text layer.]"

            if len(combined_text) == 0 and not result["is_scanned"]:
                result["error"] = "No readable text found in document."
                return result

            result["success"] = True
            return result

        except Exception as e:
            result["is_corrupt"] = True
            result["error"] = f"Error reading PDF content: {e!s}"
            return result

    @staticmethod
    def extract_docx_from_bytes(file_bytes: bytes, filename: str = "document.docx") -> dict[str, Any]:
        """
        Extract text from DOCX document using python-docx if installed,
        or standard zipfile + XML element tree parser.
        """
        result = {
            "success": False,
            "text": "",
            "page_count": 1,
            "is_scanned": False,
            "is_encrypted": False,
            "is_corrupt": False,
            "error": None,
        }

        if not file_bytes or len(file_bytes) == 0:
            result["error"] = "Empty file (0 bytes)."
            result["is_corrupt"] = True
            return result

        try:
            # Try python-docx if available
            try:
                import io
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
                combined = "\n".join(paragraphs).strip()
                result["text"] = combined
                result["success"] = bool(combined)
                if not combined:
                    result["error"] = "No readable text found in DOCX."
                return result
            except ImportError:
                pass

            # Fast native fallback via zipfile and XML
            import io
            import zipfile
            import xml.etree.ElementTree as ET

            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                xml_content = z.read("word/document.xml")
                root = ET.fromstring(xml_content)
                # XML namespace for WordprocessingML
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                paragraphs = []
                for p in root.findall(".//w:p", ns):
                    texts = [node.text for node in p.findall(".//w:t", ns) if node.text]
                    if texts:
                        paragraphs.append("".join(texts))
                combined = "\n".join(paragraphs).strip()
                result["text"] = combined
                result["success"] = bool(combined)
                if not combined:
                    result["error"] = "No readable text found in DOCX."
                return result
        except Exception as e:
            result["is_corrupt"] = True
            result["error"] = f"Failed to parse DOCX document: {e!s}"
            return result

    @staticmethod
    def extract_document_from_bytes(file_bytes: bytes, filename: str) -> dict[str, Any]:
        """Auto-detect format (PDF vs DOCX) and parse appropriately."""
        lower_name = filename.lower()
        if lower_name.endswith(".docx"):
            return PDFParser.extract_docx_from_bytes(file_bytes, filename=filename)
        return PDFParser.extract_text_from_bytes(file_bytes, filename=filename)

    @staticmethod
    def extract_text_from_file(file_path: str) -> dict[str, Any]:
        """Convenience method to parse a PDF or DOCX directly from filesystem."""
        try:
            with open(file_path, "rb") as f:
                return PDFParser.extract_document_from_bytes(f.read(), filename=file_path)
        except Exception as e:
            return {
                "success": False,
                "text": "",
                "page_count": 0,
                "is_scanned": False,
                "is_encrypted": False,
                "is_corrupt": True,
                "error": f"File access error: {e!s}",
            }
